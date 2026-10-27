"""
Ad-hoc Mention Stream Grid Generator
=====================================
Pulls ONE Cision One mention stream and writes it to a single-sheet workbook in
the legacy Meltwater / IC Data column format. Reach (and blank Title / Hit
Sentence) are enriched from CisionOne website CSV exports exactly like
generate_ic_data_grid.py does.

Unlike generate_ic_data_grid.py there is NO filtering, routing, LLM review,
finance-source rerouting, person splitting, or tier coloring: every mention the
API returns for the stream and date range ends up as one row, for manual review.

Usage:
    export CISION_API_TOKEN="your-token-here"
    python generate_adhoc_grid.py --stream-id 941746 --after 2026-05-01 --before 2026-06-01

    # Custom Input Name / output path / exports folder:
    python generate_adhoc_grid.py --stream-id 941746 --after 2026-05-01 --before 2026-06-01 \
        --input-name "My Ad-hoc Stream" --output my_pull.xlsx --cision-exports-dir cision_exports

Requires generate_ic_data_grid.py in the same folder (its API client, field
mapping, and CSV-export enrichment are reused as-is so the output stays identical
in format).
"""

import argparse
import os
import sys
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter

from generate_ic_data_grid import (
    IC_COLUMNS,
    COL_WIDTHS,
    CisionOneClient,
    mention_to_row,
    load_cision_exports,
    enrich_row_from_csv_export,
    _write_data_cell,
    HEADER_FONT,
    HEADER_FILL,
    HEADER_ALIGN,
    NO_FILL,
)

# Legacy format = the original 42 columns; the trailing LLM-review columns
# (Review Evidence/Classification/Justification/Confidence) are not used here.
ADHOC_COLUMNS = IC_COLUMNS[: IC_COLUMNS.index("Review Evidence")]


def write_sheet(ws, rows: list) -> None:
    for col_idx, header in enumerate(ADHOC_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN

    rows_sorted = sorted(
        rows,
        key=lambda r: (r[0] or datetime.min.date(), r[1] or datetime.min.time()),
        reverse=True,
    )
    for row_idx, row_data in enumerate(rows_sorted, start=2):
        for col_idx, value in enumerate(row_data[: len(ADHOC_COLUMNS)], start=1):
            _write_data_cell(ws, row_idx, col_idx, value, NO_FILL)

    for col_letter, width in COL_WIDTHS.items():
        if openpyxl.utils.column_index_from_string(col_letter) <= len(ADHOC_COLUMNS):
            ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(ADHOC_COLUMNS))}{len(rows) + 1}"


def main():
    parser = argparse.ArgumentParser(description="Generate an ad-hoc single-stream grid from the Cision One API")
    parser.add_argument("--stream-id", required=True, type=int, help="Cision One mention stream ID")
    parser.add_argument("--after", required=True, help="Start date (YYYY-MM-DD)")
    parser.add_argument("--before", required=True, help="End date (YYYY-MM-DD)")
    parser.add_argument("--input-name", default=None, help="Value for the Input Name column (default: 'Stream <id>')")
    parser.add_argument("--sheet-name", default="Data", help="Worksheet name (default: Data)")
    parser.add_argument("--output", default=None, help="Output path (default: Adhoc_Grid_<id>_<after>_to_<before>.xlsx)")
    parser.add_argument("--token", default=None, help="Cision One API token (or set CISION_API_TOKEN env var)")
    parser.add_argument(
        "--cision-exports-dir", default="cision_exports",
        help="Folder of CisionOne website CSV exports to join on Document ID for Reach etc. "
             "(default: cision_exports/; missing folder just skips enrichment)",
    )
    args = parser.parse_args()

    token = args.token or os.environ.get("CISION_API_TOKEN")
    if not token:
        print("Error: No API token. Pass --token or set CISION_API_TOKEN.")
        sys.exit(1)

    input_name = args.input_name or f"Stream {args.stream_id}"
    output_path = args.output or f"Adhoc_Grid_{args.stream_id}_{args.after}_to_{args.before}.xlsx"
    after_iso = f"{args.after}T00:00:00.000Z"
    before_iso = f"{args.before}T23:59:59.000Z"

    # mention_to_row only needs these keys; "Non-T1" keeps Custom Categories blank,
    # and a tab name outside NO_GODADDY_FOCUS_TABS just sets an unused reporting flag.
    stream_config = {"tier": "Non-T1", "tab": args.sheet_name, "input_name": input_name}

    csv_lookup = load_cision_exports(args.cision_exports_dir)

    print(f"Pulling stream {args.stream_id} ({input_name}) for {args.after} → {args.before}")
    client = CisionOneClient(api_token=token)
    mentions = client.get_all_mentions_chunked(args.stream_id, after_iso, before_iso, chunk_days=1)
    print(f"  ✓ {len(mentions)} mentions")

    rows = [mention_to_row(m, stream_config) for m in mentions]
    for row in rows:
        enrich_row_from_csv_export(row, csv_lookup)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = args.sheet_name[:31]
    write_sheet(ws, rows)
    wb.save(output_path)
    print(f"✅ Wrote {len(rows)} rows → {output_path}")


if __name__ == "__main__":
    main()
